"""Convert the Markdown reports into Word documents.

    python scripts/build_report.py

reports/REPORT.md              -> reports/MLOps_Assignment1_Report.docx (main, ~10 pages)
reports/SUPPORTING_DOCUMENT.md -> reports/MLOps_Assignment1_Supporting_Document.docx

Supports the Markdown subset used in the reports: headings, paragraphs, **bold**,
`code`, bullet/numbered lists, tables, fenced code blocks and images. Image extensions:
``![caption](path){width=5.2 crop=0.3-0.45}`` sets the width in inches and keeps only a
vertical band of the picture; several images on one line are placed side by side.
Images that do not exist are replaced by a visible placeholder.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
DOCUMENTS = {
    REPORTS / "REPORT.md": REPORTS / "MLOps_Assignment1_Report.docx",
    REPORTS / "SUPPORTING_DOCUMENT.md": REPORTS / "MLOps_Assignment1_Supporting_Document.docx",
}
TEXT_WIDTH = 7.0  # Letter page with 0.75" margins
DEFAULT_IMAGE_WIDTH = 6.3

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
IMAGE = re.compile(r"!\[(?P<alt>[^\]]*)\]\((?P<path>[^)]+)\)(?:\{(?P<attrs>[^}]*)\})?")
IMAGE_LINE = re.compile(rf"^(?:\s*{IMAGE.pattern})+\s*$")


def add_inline(paragraph, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif part.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            run = paragraph.add_run(
                label if url.startswith(("#", ".", "/")) else f"{label} ({url})"
            )
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        else:
            paragraph.add_run(part)


def _parse_attrs(attrs: str | None) -> dict[str, str]:
    return dict(item.split("=", 1) for item in (attrs or "").split() if "=" in item)


def _picture(path: Path, crop: str | None):
    if not crop:
        return str(path)
    top, bottom = (float(v) for v in crop.split("-"))
    with Image.open(path) as img:
        width, height = img.size
        band = img.crop((0, int(top * height), width, int(bottom * height)))
        buf = io.BytesIO()
        band.save(buf, format="PNG")
    buf.seek(0)
    return buf


def _caption(paragraph, alt: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(alt)
    run.italic = True
    run.font.size = Pt(9)


def _place_image(paragraph, base: Path, rel_path: str, attrs: dict, width: float) -> None:
    path = (base / rel_path).resolve()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if path.exists():
        width = float(attrs.get("width", width))
        paragraph.add_run().add_picture(_picture(path, attrs.get("crop")), width=Inches(width))
    else:
        run = paragraph.add_run(f"[Screenshot pending: {rel_path}]")
        run.italic = True
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def add_images(doc: Document, base: Path, line: str) -> None:
    images = list(IMAGE.finditer(line))
    if len(images) == 1:
        m = images[0]
        _place_image(
            doc.add_paragraph(), base, m["path"], _parse_attrs(m["attrs"]), DEFAULT_IMAGE_WIDTH
        )
        if m["alt"]:
            _caption(doc.add_paragraph(), m["alt"])
        return
    table = doc.add_table(rows=1, cols=len(images))
    cell_width = TEXT_WIDTH / len(images) - 0.15
    for cell, m in zip(table.rows[0].cells, images, strict=True):
        _place_image(cell.paragraphs[0], base, m["path"], _parse_attrs(m["attrs"]), cell_width)
        if m["alt"]:
            _caption(cell.add_paragraph(), m["alt"])


def add_table(doc: Document, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Light Grid Accent 1"
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for i, row in enumerate(cells):
        for j, value in enumerate(row[: len(cells[0])]):
            cell = table.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], value)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9)
                run.bold = run.bold or i == 0
            # keep the table on one page
            cell.paragraphs[0].paragraph_format.keep_with_next = i < len(cells) - 1
    doc.add_paragraph()


def build(src: Path, out: Path) -> Path:
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Inches(0.75)
        section.top_margin = section.bottom_margin = Inches(0.75)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(4)

    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith("```"):
            lang = stripped[3:].strip()
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            if lang != "mermaid":  # mermaid is rendered as a PNG elsewhere in the report
                p = doc.add_paragraph()
                run = p.add_run("\n".join(block))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
        elif stripped.startswith("<!--"):
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
        elif stripped == "\\pagebreak":
            doc.add_page_break()
        elif m := re.match(r"^(#{1,4})\s+(.*)", stripped):
            level = len(m.group(1))
            if level == 1 and not doc.paragraphs:
                doc.add_heading(m.group(2), level=0)
            else:
                doc.add_heading(m.group(2), level=min(level, 3) if level > 1 else 1)
        elif IMAGE_LINE.match(stripped):
            add_images(doc, src.parent, stripped)
        elif stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^\s*[-*]\s+", line):
            indent = (len(line) - len(line.lstrip())) // 2
            p = doc.add_paragraph(style="List Bullet 2" if indent else "List Bullet")
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
        elif m := re.match(r"^\s*(\d+)\.\s+(.*)", line):
            # Word's "List Number" style never restarts, so number explicitly
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.2)
            add_inline(p, f"{m[1]}. {m[2]}")
        elif stripped.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            add_inline(p, stripped.lstrip("> "))
        elif stripped:
            para = stripped
            while (
                i + 1 < len(lines)
                and lines[i + 1].strip()
                and not re.match(r"^(#|!\[|\||```|\s*[-*]\s|\s*\d+\.\s|>)", lines[i + 1].strip())
            ):
                # two trailing spaces = Markdown hard line break
                para += "\n" if lines[i].endswith("  ") else " "
                i += 1
                para += lines[i].strip()
            add_inline(doc.add_paragraph(), para)
        i += 1

    doc.save(out)
    print(f"Saved {out}")
    return out


if __name__ == "__main__":
    for source, target in DOCUMENTS.items():
        build(source, target)
